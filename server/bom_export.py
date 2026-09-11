"""Build a formatted DigiSearch Excel workbook: BOM + Part Details sheets."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
import re
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo

# Warm white / copper — paper-first, heat in the headers and totals.
PAPER = "FFFDF9"
CREAM = "F8F1E6"
SAND = "F0E4D2"
COPPER = "B5683A"
COPPER_DEEP = "8C4A28"
INK = "2C241C"
MUTED = "6B5C4E"
WHITE = "FFFFFF"
LINE = "E2D3C0"
GOOD = "3D6B45"
WARN = "9A6B1F"
BAD = "A33B2B"

THIN = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)
RULE = Border(bottom=Side(style="medium", color=COPPER))

FILL_PAPER = PatternFill("solid", fgColor=PAPER)
FILL_SAND = PatternFill("solid", fgColor=SAND)
FILL_COPPER = PatternFill("solid", fgColor=COPPER)
FILL_HEAD = PatternFill("solid", fgColor=COPPER_DEEP)
FILL_BAND = PatternFill("solid", fgColor=CREAM)
FILL_TOTAL = PatternFill("solid", fgColor=SAND)
FILL_GOOD = PatternFill("solid", fgColor="E7F0E6")
FILL_WARN = PatternFill("solid", fgColor="F7ECD4")
FILL_BAD = PatternFill("solid", fgColor="F6E0DA")

FONT_MARK = Font(name="Calibri", size=9, bold=True, color=WHITE)
FONT_TITLE = Font(name="Calibri", size=20, bold=True, color=COPPER_DEEP)
FONT_SUB = Font(name="Calibri", size=12, bold=True, color=INK)
FONT_META = Font(name="Calibri", size=10, color=MUTED)
FONT_HEAD = Font(name="Calibri", size=10, bold=True, color=WHITE)
FONT_BODY = Font(name="Calibri", size=10, color=INK)
FONT_MONO = Font(name="Consolas", size=10, color=INK)
FONT_LINK = Font(name="Calibri", size=10, color=COPPER, underline="single")
FONT_TOTAL = Font(name="Calibri", size=10, bold=True, color=COPPER_DEEP)
FONT_NOTE = Font(name="Calibri", size=9, italic=True, color=MUTED)

CENTER = Alignment(horizontal="center", vertical="center")
CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")
LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")

# Procurement-only: identity + qty + money. Catalog fields live on Part Details.
# Excel table column names cannot contain / \ * ? : [ ]
BOM_HEADERS = [
    "Item", "Designator", "Qty", "Manufacturer", "MPN", "DigiKey PN",
    "Description", "Unit Price", "Ext Price",
]
BOM_WIDTHS = [8, 16, 10, 22, 24, 28, 44, 13, 14]

DETAIL_HEADERS = [
    "Item", "Designator", "Qty", "Manufacturer", "MPN", "DigiKey PN",
    "Description", "Category", "Series", "Specs", "Stock", "Status",
    "Unit Price", "MOQ", "Price at 100", "Lifecycle", "RoHS", "REACH", "MSL",
    "ECCN", "HTSUS", "Lead weeks", "Discontinued", "End of life", "NCNR",
    "Marketplace", "Tariff active", "Match", "Other names", "Price breaks",
    "Datasheet", "Product page", "Photo",
]
DETAIL_WIDTHS = [
    7, 16, 10, 22, 24, 28, 38, 22, 16, 40, 12, 10,
    13, 10, 14, 14, 10, 18, 12,
    12, 14, 12, 14, 14, 10,
    14, 14, 28, 32, 28,
    18, 18, 14,
]

# Search results: same catalog fields as Part Details, minus BOM-only Designator/Qty.
RESULTS_HEADERS = [
    "Item", "Manufacturer", "MPN", "DigiKey PN",
    "Description", "Category", "Series", "Specs", "Stock", "Status",
    "Unit Price", "MOQ", "Price at 100", "Lifecycle", "RoHS", "REACH", "MSL",
    "ECCN", "HTSUS", "Lead weeks", "Discontinued", "End of life", "NCNR",
    "Marketplace", "Tariff active", "Match", "Other names", "Price breaks",
    "Datasheet", "Product page", "Photo",
]
RESULTS_WIDTHS = [
    7, 22, 24, 28, 38, 22, 16, 40, 12, 10,
    13, 10, 14, 14, 10, 18, 12,
    12, 14, 12, 14, 14, 10,
    14, 14, 28, 32, 28,
    18, 18, 14,
]


def _cell(ws, row, col, value, *, font=FONT_BODY, fill=None, align=LEFT, num_fmt=None, border=True):
    cell = ws.cell(row, col)
    cell.value = "" if value is None else value
    cell.font = font
    cell.alignment = align
    cell.fill = fill if fill is not None else FILL_PAPER
    if border:
        cell.border = THIN
    if num_fmt:
        cell.number_format = num_fmt
    return cell


def _banner_line(ws, row, cols, value, font, fill, height, *, rule=False):
    ws.row_dimensions[row].height = height
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    cell = ws.cell(row, 1, value)
    cell.font = font
    cell.fill = fill
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    if rule:
        cell.border = RULE
    return cell


def _banner(ws, cols: int, title: str, subtitle: str, meta: str) -> None:
    _banner_line(ws, 1, cols, "DigiSearch", FONT_MARK, FILL_COPPER, 18)
    _banner_line(ws, 2, cols, title, FONT_TITLE, FILL_PAPER, 28)
    _banner_line(ws, 3, cols, subtitle, FONT_SUB, FILL_PAPER, 18)
    _banner_line(ws, 4, cols, meta, FONT_META, FILL_PAPER, 18, rule=True)


def _add_table(ws, name: str, headers: list[str], header_row: int, last_row: int, *, last_column=False) -> None:
    columns = [TableColumn(id=i, name=header) for i, header in enumerate(headers, start=1)]
    last_col = get_column_letter(len(headers))
    ref = f"A{header_row}:{last_col}{last_row}"
    table = Table(
        displayName=name,
        name=name,
        ref=ref,
        headerRowCount=1,
        totalsRowCount=0,
        tableColumns=columns,
        autoFilter=AutoFilter(ref=ref),
        tableStyleInfo=TableStyleInfo(
            name="TableStyleLight8",
            showFirstColumn=False,
            showLastColumn=last_column,
            showRowStripes=True,
            showColumnStripes=False,
        ),
    )
    ws.add_table(table)


def _print_setup(ws, title_rows: str, footer: str, *, tabloid: bool = False) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_TABLOID if tabloid else ws.PAPERSIZE_LETTER
    ws.page_margins.left = 0.5
    ws.page_margins.right = 0.5
    ws.page_margins.top = 0.6
    ws.page_margins.bottom = 0.5
    ws.oddHeader.left.text = "DigiSearch"
    ws.oddFooter.left.text = footer
    ws.oddFooter.right.text = "Page &P of &N"
    ws.print_title_rows = title_rows
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _part(line: dict) -> dict:
    part = line.get("part")
    return part if isinstance(part, dict) else {}


def _breaks(part: dict) -> str:
    breaks = part.get("priceBreaks") or []
    bits = []
    for b in breaks:
        try:
            bits.append(f"{int(b.get('qty', 0))}: ${float(b.get('price', 0)):.4g}")
        except (TypeError, ValueError):
            continue
    return " · ".join(bits)


def _yes_no(value: Any) -> str:
    if value is True:
        return "Yes"
    if value is False:
        return "No"
    return ""


def _sanitize_xlsx(data: bytes) -> bytes:
    """Drop empty cached formula values that make Excel show a repair prompt."""
    src = ZipFile(BytesIO(data), "r")
    out = BytesIO()
    with ZipFile(out, "w", ZIP_DEFLATED) as dst:
        for name in src.namelist():
            payload = src.read(name)
            if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
                payload = re.sub(rb"<f>([^<]*)</f><v></v>", rb"<f>\1</f>", payload)
            dst.writestr(name, payload)
    src.close()
    return out.getvalue()


def build_bom_workbook(lines: list[dict]) -> bytes:
    now = datetime.now().astimezone()
    stamp = now.strftime("%Y-%m-%d %H:%M")
    units = sum(int(l.get("qty") or 0) for l in lines)
    cost = sum(float(l.get("price") or 0) * int(l.get("qty") or 0) for l in lines)

    wb = Workbook()
    wb.properties.title = "DigiSearch Bill of Materials"
    wb.properties.creator = "DigiSearch"
    _write_bom_sheet(wb.active, lines, stamp, units, cost)
    _write_details_sheet(wb.create_sheet("Part Details"), lines, stamp)

    buf = BytesIO()
    wb.save(buf)
    return _sanitize_xlsx(buf.getvalue())


def _write_bom_sheet(ws, lines: list[dict], stamp: str, units: int, cost: float) -> None:
    ws.title = "BOM"
    ws.sheet_properties.tabColor = COPPER
    cols = len(BOM_HEADERS)
    _banner(
        ws, cols,
        "Bill of Materials",
        "Quantity, unit price, and extended cost",
        f"Exported {stamp}    {len(lines)} line item{'s' if len(lines) != 1 else ''}    "
        f"{units} unit{'s' if units != 1 else ''}    Estimated {cost:,.2f} USD",
    )

    header_row = 6
    first_data = header_row + 1
    last_data = header_row + max(len(lines), 1)
    total_row = last_data + 1

    for i, (title, width) in enumerate(zip(BOM_HEADERS, BOM_WIDTHS), start=1):
        _cell(ws, header_row, i, title, font=FONT_HEAD, fill=FILL_HEAD, align=CENTER)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[header_row].height = 20

    data_lines = lines or [{}]
    for idx, line in enumerate(data_lines, start=1):
        row = header_row + idx
        part = _part(line)
        fill = FILL_BAND if idx % 2 == 0 else FILL_PAPER
        qty = int(line.get("qty") or 0)
        price = float(line.get("price") or 0)
        values = [
            (idx if lines else "", FONT_BODY, CENTER, "0"),
            (line.get("desig") or "", FONT_MONO, LEFT, None),
            (qty if lines else "", FONT_BODY, CENTER, "0"),
            (line.get("mfr") or "", FONT_BODY, LEFT, None),
            (line.get("mpn") or "", FONT_MONO, LEFT, None),
            (line.get("dk") or "", FONT_MONO, LEFT, None),
            (part.get("description") or "", FONT_BODY, LEFT_WRAP, None),
            (price if lines else "", FONT_MONO, RIGHT, '"$"#,##0.0000'),
            ((qty * price) if lines else "", FONT_MONO, RIGHT, '"$"#,##0.00'),
        ]
        ws.row_dimensions[row].height = 20
        for col, (val, font, align, fmt) in enumerate(values, start=1):
            _cell(ws, row, col, val, font=font, fill=fill, align=align, num_fmt=fmt)

    _add_table(ws, "BOMLines", BOM_HEADERS, header_row, last_data, last_column=True)

    ws.row_dimensions[total_row].height = 20
    for col in range(1, cols + 1):
        if col in (1, 3, 9):
            continue
        _cell(ws, total_row, col, "", font=FONT_TOTAL, fill=FILL_TOTAL, align=LEFT)
    _cell(ws, total_row, 1, "Total", font=FONT_TOTAL, fill=FILL_TOTAL, align=LEFT)
    _cell(
        ws, total_row, 3,
        f"=SUM(C{first_data}:C{last_data})" if lines else 0,
        font=FONT_TOTAL, fill=FILL_TOTAL, align=CENTER, num_fmt="0",
    )
    _cell(
        ws, total_row, 9,
        f"=SUM(I{first_data}:I{last_data})" if lines else 0,
        font=FONT_TOTAL, fill=FILL_TOTAL, align=RIGHT, num_fmt='"$"#,##0.00',
    )

    note_row = total_row + 2
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=cols)
    note = ws.cell(note_row, 1, "Specs, stock, lifecycle, datasheets, and compliance fields are on the Part Details sheet.")
    note.font = FONT_NOTE
    note.fill = FILL_PAPER
    note.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws.freeze_panes = f"A{first_data}"
    ws.print_area = f"A1:{get_column_letter(cols)}{note_row}"
    _print_setup(ws, "1:5", "Bill of Materials")


def _write_details_sheet(ws, lines: list[dict], stamp: str) -> None:
    ws.title = "Part Details"
    ws.sheet_properties.tabColor = SAND
    cols = len(DETAIL_HEADERS)
    _banner(
        ws, cols,
        "Part Details",
        "Catalog, stock, lifecycle, and sourcing fields for every BOM line",
        f"Full listing data  ·  {stamp}",
    )

    header_row = 6
    first_data = header_row + 1
    last_data = header_row + max(len(lines), 1)

    for i, (title, width) in enumerate(zip(DETAIL_HEADERS, DETAIL_WIDTHS), start=1):
        _cell(ws, header_row, i, title, font=FONT_HEAD, fill=FILL_HEAD, align=CENTER_WRAP)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[header_row].height = 28

    data_lines = lines or [{}]
    for idx, line in enumerate(data_lines, start=1):
        row = header_row + idx
        part = _part(line)
        fill = FILL_BAND if idx % 2 == 0 else FILL_PAPER
        qty = int(line.get("qty") or 0)
        price = float(line.get("price") or 0)
        datasheet = (part.get("datasheetUrl") or "").strip()
        product = (part.get("productUrl") or "").strip()
        photo = (part.get("photoUrl") or "").strip()
        match_total = part.get("matchTotal")
        match_score = part.get("matchScore")
        match = ""
        if match_score is not None and match_total is not None:
            match = f"{match_score} / {match_total}"
            note = part.get("matchNote") or ""
            if note:
                match = f"{match} — {note}"
        status = (line.get("status") or part.get("status") or "").lower()
        stock_fill = FILL_BAD if status == "out" else FILL_WARN if status == "low" else FILL_GOOD if status == "in" else fill

        values = [
            idx if lines else "", line.get("desig") or "", qty if lines else "",
            line.get("mfr") or part.get("mfr") or "",
            line.get("mpn") or part.get("mpn") or "", line.get("dk") or part.get("dk") or "",
            part.get("description") or "", part.get("category") or "", part.get("series") or "",
            part.get("attrs") or "", int(line.get("stock") or part.get("stock") or 0) if lines else "",
            line.get("status") or part.get("status") or "",
            price if lines else "", int(part.get("moq") or 1) if lines else "", part.get("priceAt100"),
            part.get("lifecycle") or "",
            "Yes" if part.get("rohs") else "No" if "rohs" in part else "",
            part.get("reachStatus") or "", part.get("moistureSensitivityLevel") or "",
            part.get("exportControlClassNumber") or "", part.get("htsusCode") or "",
            part.get("leadWeeks") if part.get("leadWeeks") not in (None, "") else "",
            _yes_no(part.get("discontinued")), _yes_no(part.get("endOfLife")),
            _yes_no(part.get("ncnr")), _yes_no(part.get("marketplace")),
            _yes_no(part.get("tariffActive")),
            match, ", ".join(part.get("otherNames") or []), _breaks(part),
            datasheet or "—", product or "—", photo or "—",
        ]
        ws.row_dimensions[row].height = 20
        for col, val in enumerate(values, start=1):
            font = FONT_MONO if col in (2, 5, 6, 20, 21) else FONT_BODY
            align = CENTER if col in (1, 3, 12, 16, 17, 23, 24, 25, 26, 27) else LEFT_WRAP if col in (7, 10, 28, 29, 30) else LEFT
            fmt = None
            cell_fill = fill
            if col == 11:
                cell_fill = stock_fill
            if col == 13:
                fmt = '"$"#,##0.0000'
                align = RIGHT
                font = FONT_MONO
            elif col == 15:
                fmt = '"$"#,##0.0000'
                align = RIGHT
                font = FONT_MONO
            elif col in (3, 11, 14):
                fmt = "#,##0"
                align = RIGHT if col != 3 else CENTER
            cell = _cell(ws, row, col, val if val is not None else "", font=font, fill=cell_fill, align=align, num_fmt=fmt)
            if col == 31 and datasheet:
                cell.value = "Open datasheet"
                cell.font = FONT_LINK
                cell.hyperlink = datasheet
            elif col == 32 and product:
                cell.value = "Open product page"
                cell.font = FONT_LINK
                cell.hyperlink = product
            elif col == 33 and photo:
                cell.value = "Open photo"
                cell.font = FONT_LINK
                cell.hyperlink = photo

    _add_table(ws, "PartDetails", DETAIL_HEADERS, header_row, last_data)
    ws.freeze_panes = f"C{first_data}"
    ws.print_area = f"A1:{get_column_letter(cols)}{last_data}"
    _print_setup(ws, "1:5", "Part Details", tabloid=True)


def build_search_results_workbook(rows: list[dict]) -> bytes:
    now = datetime.now().astimezone()
    stamp = now.strftime("%Y-%m-%d %H:%M")

    wb = Workbook()
    wb.properties.title = "DigiSearch Search Results"
    wb.properties.creator = "DigiSearch"
    _write_results_sheet(wb.active, rows, stamp)

    buf = BytesIO()
    wb.save(buf)
    return _sanitize_xlsx(buf.getvalue())


def _write_results_sheet(ws, rows: list[dict], stamp: str) -> None:
    ws.title = "Search Results"
    ws.sheet_properties.tabColor = COPPER
    cols = len(RESULTS_HEADERS)
    _banner(
        ws, cols,
        "Search Results",
        "Catalog, stock, lifecycle, and sourcing fields for every match on screen",
        f"Exported {stamp}    {len(rows)} result{'s' if len(rows) != 1 else ''}",
    )

    header_row = 6
    first_data = header_row + 1
    last_data = header_row + max(len(rows), 1)

    for i, (title, width) in enumerate(zip(RESULTS_HEADERS, RESULTS_WIDTHS), start=1):
        _cell(ws, header_row, i, title, font=FONT_HEAD, fill=FILL_HEAD, align=CENTER_WRAP)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[header_row].height = 28

    data_rows = rows or [{}]
    for idx, part in enumerate(data_rows, start=1):
        row = header_row + idx
        fill = FILL_BAND if idx % 2 == 0 else FILL_PAPER
        datasheet = (part.get("datasheetUrl") or "").strip()
        product = (part.get("productUrl") or "").strip()
        photo = (part.get("photoUrl") or "").strip()
        match_total = part.get("matchTotal")
        match_score = part.get("matchScore")
        match = ""
        if match_score is not None and match_total is not None:
            match = f"{match_score} / {match_total}"
            note = part.get("matchNote") or ""
            if note:
                match = f"{match} — {note}"
        status = (part.get("status") or "").lower()
        stock_fill = FILL_BAD if status == "out" else FILL_WARN if status == "low" else FILL_GOOD if status == "in" else fill

        values = [
            idx if rows else "",
            part.get("mfr") or "", part.get("mpn") or "", part.get("dk") or "",
            part.get("description") or "", part.get("category") or "", part.get("series") or "",
            part.get("attrs") or "", int(part.get("stock") or 0) if rows else "",
            part.get("status") or "",
            float(part.get("price") or 0) if rows else "", int(part.get("moq") or 1) if rows else "",
            part.get("priceAt100"),
            part.get("lifecycle") or "",
            "Yes" if part.get("rohs") else "No" if "rohs" in part else "",
            part.get("reachStatus") or "", part.get("moistureSensitivityLevel") or "",
            part.get("exportControlClassNumber") or "", part.get("htsusCode") or "",
            part.get("leadWeeks") if part.get("leadWeeks") not in (None, "") else "",
            _yes_no(part.get("discontinued")), _yes_no(part.get("endOfLife")),
            _yes_no(part.get("ncnr")), _yes_no(part.get("marketplace")),
            _yes_no(part.get("tariffActive")),
            match, ", ".join(part.get("otherNames") or []), _breaks(part),
            datasheet or "—", product or "—", photo or "—",
        ]
        ws.row_dimensions[row].height = 20
        for col, val in enumerate(values, start=1):
            font = FONT_MONO if col in (3, 4, 18, 19) else FONT_BODY
            align = CENTER if col in (1, 10, 14, 15, 21, 22, 23, 24, 25) else LEFT_WRAP if col in (5, 8, 26, 27, 28) else LEFT
            fmt = None
            cell_fill = fill
            if col == 9:
                cell_fill = stock_fill
            if col == 11:
                fmt = '"$"#,##0.0000'
                align = RIGHT
                font = FONT_MONO
            elif col == 13:
                fmt = '"$"#,##0.0000'
                align = RIGHT
                font = FONT_MONO
            elif col in (9, 12):
                fmt = "#,##0"
                align = RIGHT
            cell = _cell(ws, row, col, val if val is not None else "", font=font, fill=cell_fill, align=align, num_fmt=fmt)
            if col == 29 and datasheet:
                cell.value = "Open datasheet"
                cell.font = FONT_LINK
                cell.hyperlink = datasheet
            elif col == 30 and product:
                cell.value = "Open product page"
                cell.font = FONT_LINK
                cell.hyperlink = product
            elif col == 31 and photo:
                cell.value = "Open photo"
                cell.font = FONT_LINK
                cell.hyperlink = photo

    _add_table(ws, "SearchResults", RESULTS_HEADERS, header_row, last_data)
    ws.freeze_panes = f"B{first_data}"
    ws.print_area = f"A1:{get_column_letter(cols)}{last_data}"
    _print_setup(ws, "1:5", "Search Results", tabloid=True)
